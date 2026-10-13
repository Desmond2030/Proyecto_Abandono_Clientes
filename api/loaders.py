# ESTADO_ARCHIVO: EN_MODIFICACION
"""Adaptación de las clases existentes de src sin modificar el entrenamiento."""
import csv
from pathlib import Path

import joblib
import matplotlib
matplotlib.use("Agg")
from openpyxl import load_workbook

from src.Data_Processor import DataProcessor

EXTENSIONES = {"modelo": {".keras"}, "scaler": {".pkl", ".joblib"},
               "datos": {".csv", ".xlsx", ".parquet"}}


def cargar_archivo(path, tipo):
    path = Path(path)
    if tipo not in EXTENSIONES or path.suffix.lower() not in EXTENSIONES[tipo]:
        raise ValueError("El formato no corresponde al tipo de carga seleccionado.")
    if tipo == "modelo":
        from api.model_adapter import ModeloInferencia
        return ModeloInferencia().cargar_modelo(path)
    processor = DataProcessor()
    if tipo == "scaler":
        if path.suffix.lower() == ".pkl":
            processor.save_load_scaler(str(path), load=True)
        else:
            processor.scaler = joblib.load(path)
        return processor
    # pandas renombra encabezados CSV/Excel duplicados; detectarlos antes.
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            headers = next(csv.reader(stream), [])
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            headers = list(next(workbook.active.iter_rows(values_only=True), []))
        finally:
            workbook.close()
    else:
        headers = None
    if headers is not None and len(headers) != len(set(headers)):
        raise ValueError("Hay encabezados duplicados en el archivo.")
    options = {"nrows": 10001} if path.suffix.lower() in {".csv", ".xlsx"} else {}
    return processor.load_data(str(path), **options)
